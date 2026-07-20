from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/Guia_de_Instalacion_y_Respaldo_Soft_Inventario.docx")
BLUE = RGBColor(31, 78, 121)
DARK = RGBColor(32, 44, 56)
GRAY = RGBColor(90, 99, 108)
GREEN = RGBColor(30, 115, 82)
RED = RGBColor(166, 47, 47)


def set_font(run, size=11, bold=False, color=DARK, italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=120, start=160, bottom=120, end=160):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_callout(doc, title, text, tone="info"):
    palette = {
        "info": ("E8EEF5", BLUE),
        "ok": ("E8F3EE", GREEN),
        "warn": ("FCEAEA", RED),
    }
    fill, color = palette[tone]
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(title), 11, True, color)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    set_font(p.add_run(text), 10.5, False, DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_step(doc, number, title, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.2
    set_font(p.add_run(title + ": "), 11, True, DARK)
    set_font(p.add_run(text), 11, False, DARK)


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    if bold_prefix and text.startswith(bold_prefix):
        set_font(p.add_run(bold_prefix), 11, True)
        set_font(p.add_run(text[len(bold_prefix):]), 11)
    else:
        set_font(p.add_run(text), 11)


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def page_break(doc):
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.4)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = DARK
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for style_name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, DARK, 10, 5),
):
    style = styles[style_name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = color
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(header.add_run("SOFT INVENTARIO  |  GUÍA PARA EL CLIENTE"), 8.5, True, GRAY)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run("Conserve esta guía junto con su código de recuperación."), 8.5, False, GRAY)

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(80)
p.paragraph_format.space_after = Pt(6)
set_font(p.add_run("GUÍA DE USO"), 12, True, GREEN)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
set_font(p.add_run("Instalación, respaldo y recuperación"), 28, True, DARK)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(28)
set_font(p.add_run("Soft Inventario para Windows"), 15, False, BLUE)

add_callout(
    doc,
    "Esta guía está pensada para usarla paso a paso",
    "No necesita conocimientos técnicos. Lea un paso, hágalo y recién después continúe con el siguiente. Si aparece un mensaje diferente al mostrado aquí, no borre archivos: anote el mensaje o tome una foto y comuníquese con soporte.",
    "info",
)

heading(doc, "Guarde estos tres elementos", 2)
add_bullet(doc, "El instalador oficial del programa.")
add_bullet(doc, "La licencia entregada por el proveedor.")
add_bullet(doc, "El código de recuperación que mostrará el programa al crear su usuario.")

page_break(doc)

heading(doc, "1. Antes de instalar", 1)
add_bullet(doc, "Use una computadora con Windows y una cuenta que permita instalar programas.")
add_bullet(doc, "Conecte la computadora a Internet durante la instalación y activación.")
add_bullet(doc, "Cierre versiones anteriores de Soft Inventario antes de instalar una actualización.")
add_bullet(doc, "Tenga a mano el código de licencia enviado por el proveedor.")
add_callout(doc, "Importante", "Instale únicamente el archivo recibido del proveedor o descargado desde el enlace oficial. Si Windows muestra una advertencia que no reconoce, deténgase y consulte a soporte.", "warn")

heading(doc, "2. Instalar el programa", 1)
add_step(doc, 1, "Busque el instalador", "normalmente se llama Inventario.App_x.x.x_x64-setup.exe y estará en la carpeta Descargas.")
add_step(doc, 2, "Abra el instalador", "haga doble clic sobre el archivo.")
add_step(doc, 3, "Espere", "Windows instalará el programa. No apague la computadora durante este proceso.")
add_step(doc, 4, "Abra Soft Inventario", "use el icono creado en el escritorio o en el menú de aplicaciones.")
add_step(doc, 5, "No repita la instalación", "si la ventana tarda algunos segundos en abrir, espere antes de volver a hacer clic.")

heading(doc, "3. Activar la licencia", 1)
add_step(doc, 1, "Copie la licencia completa", "debe copiar todos sus caracteres, sin quitar el principio ni el final.")
add_step(doc, 2, "Pegue la licencia", "en el cuadro Token de activación RSA.")
add_step(doc, 3, "Pulse Activar Sistema", "espere hasta que aparezca la pantalla para crear el acceso.")
add_callout(doc, "Si la licencia es rechazada", "No intente modificarla. Verifique que se haya copiado completa. Si el problema continúa, envíe una foto del mensaje a soporte.", "warn")

page_break(doc)

heading(doc, "4. Crear el usuario por primera vez", 1)
add_step(doc, 1, "Escriba un nombre de usuario", "elija uno corto que pueda recordar.")
add_step(doc, 2, "Cree una contraseña", "use al menos 8 caracteres y no la comparta con personas no autorizadas.")
add_step(doc, 3, "Repita la contraseña", "debe escribir exactamente la misma contraseña.")
add_step(doc, 4, "Pulse Crear acceso", "el programa mostrará un código de recuperación.")

heading(doc, "5. Guardar el código de recuperación", 1)
add_callout(doc, "Este código es indispensable", "Permite recuperar el acceso si olvida el usuario o la contraseña. No es la licencia y no debe publicarse ni enviarse a desconocidos.", "warn")
add_bullet(doc, "Anótelo en papel y guárdelo en un lugar seguro.")
add_bullet(doc, "También puede guardarlo en un archivo protegido o imprimirlo.")
add_bullet(doc, "No lo deje únicamente dentro de la misma computadora.")
add_bullet(doc, "Pulse Ya guardé el código e ingresar solamente después de haberlo guardado.")

heading(doc, "6. Entrar y salir normalmente", 1)
add_bullet(doc, "Al abrir el programa, escriba su usuario y contraseña.")
add_bullet(doc, "Mientras el programa permanezca abierto, la sesión seguirá activa.")
add_bullet(doc, "Para cerrar, use la opción Salir o cierre la ventana normalmente.")
add_bullet(doc, "No apague la computadora mientras se esté guardando, restaurando o actualizando.")

page_break(doc)

heading(doc, "7. Crear un respaldo", 1)
p = doc.add_paragraph("Un respaldo es una copia de seguridad de los productos, variantes, stock, movimientos y configuración guardados en la base local.")

heading(doc, "Cuándo hacerlo", 2)
add_bullet(doc, "Una vez por semana.")
add_bullet(doc, "Antes de una actualización importante.")
add_bullet(doc, "Antes de cambiar de computadora.")
add_bullet(doc, "Después de cargar una gran cantidad de productos o movimientos.")

heading(doc, "Pasos para crear el respaldo", 2)
add_step(doc, 1, "Abra Configuración", "desde el menú lateral.")
add_step(doc, 2, "Busque Backup Local", "está en la parte derecha de la pantalla.")
add_step(doc, 3, "Pulse Crear Backup (.db)", "espere a que aparezca el mensaje Respaldo creado en.")
add_step(doc, 4, "Anote la ubicación", "el mensaje indica dónde se guardó el archivo.")
add_step(doc, 5, "Haga una segunda copia", "copie el archivo .db a un pendrive, disco externo o carpeta privada en la nube.")

add_callout(doc, "Regla sencilla", "Mantenga al menos dos copias: una en la computadora y otra fuera de ella. Un respaldo guardado sólo en la misma computadora no protege ante rotura, pérdida o robo.", "ok")

heading(doc, "Cómo reconocer el archivo correcto", 2)
add_bullet(doc, "El nombre contiene inventario_backup.")
add_bullet(doc, "La extensión del archivo es .db.")
add_bullet(doc, "La fecha y hora deben coincidir con el momento en que hizo el respaldo.")
add_bullet(doc, "No abra ni edite el archivo con Excel, Bloc de notas u otro programa.")

page_break(doc)

heading(doc, "8. Restaurar un respaldo", 1)
add_callout(doc, "Atención", "Restaurar reemplaza la base actual por la copia elegida. Los cambios realizados después de la fecha de ese respaldo dejarán de aparecer. El programa crea antes una copia automática de la base actual.", "warn")

add_step(doc, 1, "Cierre otras tareas", "no sincronice ni registre movimientos durante la restauración.")
add_step(doc, 2, "Abra Configuración", "busque la sección Backup Local.")
add_step(doc, 3, "Pulse Restaurar Backup", "se abrirá una ventana para elegir un archivo.")
add_step(doc, 4, "Seleccione el archivo .db", "compruebe su nombre y fecha antes de continuar.")
add_step(doc, 5, "Confirme la advertencia", "el programa validará el archivo y se reiniciará.")
add_step(doc, 6, "Espere el reinicio", "no fuerce el cierre ni apague la computadora.")
add_step(doc, 7, "Revise la información", "confirme productos, stock y movimientos recientes.")

heading(doc, "Si eligió un archivo incorrecto", 2)
add_bullet(doc, "No cree ni modifique productos hasta resolverlo.")
add_bullet(doc, "Comuníquese con soporte e indique la fecha y hora de la restauración.")
add_bullet(doc, "No borre las copias automáticas ni la carpeta de datos del programa.")

page_break(doc)

heading(doc, "9. Instalar actualizaciones", 1)
add_step(doc, 1, "Cree un respaldo", "hágalo antes de actualizar.")
add_step(doc, 2, "Abra Configuración", "busque la tarjeta Actualizaciones.")
add_step(doc, 3, "Pulse Buscar actualizaciones", "espere la respuesta del programa.")
add_step(doc, 4, "Pulse Instalar versión", "si existe una versión nueva.")
add_step(doc, 5, "Espere la descarga", "no cierre el programa durante este paso.")
add_step(doc, 6, "Permita el reinicio", "la actualización conserva la base de datos y la configuración.")
add_step(doc, 7, "Compruebe la versión", "vuelva a Configuración y verifique el número instalado.")

add_callout(doc, "No desinstale para actualizar", "Normalmente debe instalar la nueva versión sobre la existente o usar el botón de actualización. Desinstalar o borrar carpetas manualmente puede dificultar la recuperación de datos.", "warn")

heading(doc, "10. Sincronizar con Tiendanube", 1)
add_bullet(doc, "Antes de la primera importación, cree un respaldo.")
add_bullet(doc, "Use Importar desde Tiendanube para traer productos al programa.")
add_bullet(doc, "Espere hasta que el registro indique que terminó.")
add_bullet(doc, "No cierre la ventana ni desconecte Internet durante la sincronización.")
add_bullet(doc, "Si aparece ERROR, tome una foto o copie el mensaje completo y envíelo a soporte.")
add_bullet(doc, "No repita muchas veces la importación mientras exista un error sin revisar.")

page_break(doc)

heading(doc, "11. Si olvidó el usuario o la contraseña", 1)
add_step(doc, 1, "Pulse Olvidé mi contraseña", "en la pantalla de ingreso.")
add_step(doc, 2, "Ingrese el código de recuperación", "use el código guardado durante el primer acceso.")
add_step(doc, 3, "Cree las nuevas credenciales", "anote el nuevo usuario y contraseña en un lugar seguro.")
add_step(doc, 4, "Ingrese nuevamente", "compruebe que el acceso funciona.")

heading(doc, "Si perdió también el código", 2)
add_bullet(doc, "Pulse Recuperación con soporte.")
add_bullet(doc, "Envíe a soporte la solicitud generada por esa instalación.")
add_bullet(doc, "Use únicamente la respuesta firmada que entregue el soporte autorizado.")
add_callout(doc, "Seguridad", "No existe una contraseña universal del desarrollador. El soporte no necesita conocer su contraseña y nunca debe pedirle que la publique.", "ok")

heading(doc, "12. Qué nunca debe hacer", 1)
add_bullet(doc, "No borre la carpeta com.softinventario.app.")
add_bullet(doc, "No edite archivos .db, license.key ni archivos internos.")
add_bullet(doc, "No comparta su contraseña, licencia o código de recuperación con desconocidos.")
add_bullet(doc, "No use limpiadores de sistema sobre las carpetas del programa sin consultar.")
add_bullet(doc, "No apague la computadora durante una actualización, respaldo, restauración o sincronización.")

page_break(doc)

heading(doc, "13. Información para pedir soporte", 1)
p = doc.add_paragraph("Antes de comunicarse, reúna estos datos. Esto permite resolver el problema con mayor rapidez:")
add_bullet(doc, "Versión instalada, visible en Configuración.")
add_bullet(doc, "Qué acción estaba realizando.")
add_bullet(doc, "Texto completo del error o una fotografía clara.")
add_bullet(doc, "Fecha y hora aproximada del problema.")
add_bullet(doc, "Si el problema ocurrió durante Tiendanube, indique si era importación o envío de stock.")

heading(doc, "Lista rápida de control", 2)
check_table = doc.add_table(rows=1, cols=2)
check_table.autofit = False
check_table.columns[0].width = Inches(0.55)
check_table.columns[1].width = Inches(5.95)
for cell in check_table.rows[0].cells:
    shade(cell, "E8EEF5")
    set_cell_margins(cell)
set_font(check_table.cell(0, 0).paragraphs[0].add_run("OK"), 10, True, BLUE)
set_font(check_table.cell(0, 1).paragraphs[0].add_run("Antes de entregar o comenzar a trabajar"), 10, True, BLUE)
for item in (
    "El programa abre y permite iniciar sesión.",
    "La licencia aparece como validada.",
    "El código de recuperación está guardado fuera de la computadora.",
    "Existe al menos un respaldo .db fuera de la computadora.",
    "La versión instalada es la última disponible.",
    "La sincronización con Tiendanube finaliza sin errores.",
):
    row = check_table.add_row().cells
    for cell in row:
        set_cell_margins(cell)
    p = row[0].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("☐"), 13, False, DARK)
    set_font(row[1].paragraphs[0].add_run(item), 10.5, False, DARK)

doc.add_paragraph()
add_callout(doc, "Consejo final", "Cuando tenga dudas, deténgase antes de borrar o reinstalar. La mayoría de los problemas se resuelven conservando la base y revisando el mensaje de error.", "info")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT.resolve())
